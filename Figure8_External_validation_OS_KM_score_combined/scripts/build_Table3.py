"""Build the reproducibility Word table from the multi-cohort summary CSV."""

import csv
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

TABLES_DIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "results", "tables",
)
SRC = os.path.join(TABLES_DIR, "NCC_multi_cohort_validation_summary.csv")
OUT = os.path.join(TABLES_DIR, "Table3_reproducibility.docx")

with open(SRC, newline="", encoding="utf-8") as f:
    reader = csv.reader(f)
    header = next(reader)
    rows = [row for row in reader if row]

COHORTS = ["E-MTAB-1980", "ICGC RECA-EU", "CPTAC"]
COLIDX = {h: i for i, h in enumerate(header)}


def get(section, item, cohort):
    ci = COLIDX[cohort]
    for row in rows:
        if row[0] == section and row[1] == item:
            return row[ci]
    return None


def slash(v):
    return v.strip().replace(" of ", "/")


data = []
for cohort in COHORTS:
    available = slash(get("Classifier robustness", "NCC genes available", cohort))
    r_class = get("Classifier robustness", "68-gene Pearson r vs TCGA", cohort).lstrip("+")
    ci_class = get("Classifier robustness", "68-gene 95% CI", cohort)
    dir14 = slash(get("Biological validation", "14-gene direction concordance", cohort))
    r_eff = get("Biological validation", "14-gene Pearson r vs TCGA", cohort).lstrip("+")
    ci_eff = get("Biological validation", "14-gene 95% CI", cohort)

    classifier = f"{r_class} (95% CI: {ci_class})"
    effect = f"{r_eff} (95% CI: {ci_eff})"

    if cohort == "CPTAC":
        transcriptomic = "Not available"
        proteomic = f"{dir14} proteins"
    else:
        transcriptomic = f"{dir14} genes"
        proteomic = "Not available"

    data.append([cohort, available, classifier, transcriptomic, proteomic, effect])

doc = Document()
doc.styles["Normal"].font.name = "Arial"
doc.styles["Normal"].font.size = Pt(9)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = title.add_run(
    "Table 3. Cross-cohort reproducibility of lipid-metabolic subtype features"
)
run.bold = True
run.font.name = "Arial"
run.font.size = Pt(10)

hdr_texts = [
    "Validation cohort",
    "Available classifier genes",
    "Subtype-associated gene expression correlation (Pearson r)",
    "Lipid-metabolic transcriptomic concordance",
    "Lipid-metabolic proteomic concordance",
    "Effect-size correlation with TCGA-KIRC",
]

table = doc.add_table(rows=1, cols=len(hdr_texts))
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = table.rows[0].cells
for i, t in enumerate(hdr_texts):
    hdr[i].text = t


def style_cell_text(cell, bold=False, size=9, align=None):
    for p in cell.paragraphs:
        if align is not None:
            p.alignment = align
        for rn in p.runs:
            rn.font.name = "Arial"
            rn.font.size = Pt(size)
            rn.font.bold = bold


for c in hdr:
    style_cell_text(c, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    tcPr = c._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "D9E2F3")
    tcPr.append(shd)

for row in data:
    cells = table.add_row().cells
    for i, val in enumerate(row):
        cells[i].text = val
        align = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
        style_cell_text(cells[i], size=9, align=align)

notes = [
    "Available classifier genes = number of the 68 NCC classifier genes detectable in "
    "each cohort's expression matrix. E-MTAB-1980 lacked AR and SLC22A2; CPTAC lacked "
    "MTARC2, ECRG4, and CAVIN2.",
    "Pearson correlation coefficients were calculated using subtype-associated log2 "
    "fold-change (S1 vs S2) of the 68-gene classifier (gene expression correlation "
    "column) or the 14-gene lipid-metabolic panel (effect-size correlation column) "
    "between TCGA-KIRC and each validation cohort.",
    "Directional concordance indicates the proportion of lipid-metabolism-related "
    "markers showing the same direction of subtype association (S1 vs S2) as observed "
    "in TCGA-KIRC.",
    "Abbreviations: NCC, nearest-centroid classifier; TCGA-KIRC, The Cancer Genome "
    "Atlas kidney renal clear cell carcinoma; ICGC RECA-EU, International Cancer "
    "Genome Consortium renal cell carcinoma European cohort; CPTAC, Clinical "
    "Proteomic Tumor Analysis Consortium.",
    "Not available, modality not assessed in that cohort.",
]
for n in notes:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    rn = p.add_run(n)
    rn.font.name = "Arial"
    rn.font.size = Pt(8)

doc.save(OUT)
print(f"saved: {OUT}")
for row in data:
    print(row)