# table_s2_ncc_classifier.py
# Build Table S2 (NCC classifier, 68 genes) as a formatted Excel file.
# Reads data/subtype_classifier.json.

import sys, os, io, atexit, json
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment


def _wait_on_exit():
    try:
        if os.name == 'nt':
            import msvcrt
            msvcrt.getch()
    except Exception:
        pass


atexit.register(_wait_on_exit)

if sys.platform == 'win32' and sys.stdout is not None:
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8',
                                  errors='replace')

WORK = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(WORK)
RESULTS = os.path.join(REPO, 'results')
os.makedirs(RESULTS, exist_ok=True)

with open(os.path.join(REPO, 'data', 'subtype_classifier.json')) as f:
    clf = json.load(f)

rows = []
for i, g in enumerate(clf['genes']):
    rows.append({
        'Gene': g,
        'S1_Zscore': round(clf['centroid_s1'][i], 6),
        'S2_Zscore': round(clf['centroid_s2'][i], 6),
        'Training_Mean_log2TPM': round(clf['scaler_mean'][i], 6),
        'Training_SD_log2TPM': round(clf['scaler_std'][i], 6),
    })
df = pd.DataFrame(rows)

wb = Workbook()
ws = wb.active
ws.title = 'NCC Classifier'

ws.merge_cells('A1:E1')
ws['A1'] = ('Supplementary Table S2. Nearest-centroid classifier (NCC) '
            'based on 68 lipid metabolism-related genes.')
ws['A1'].font = Font(bold=True, size=12)
ws['A1'].alignment = Alignment(wrap_text=True)
ws.row_dimensions[1].height = 25

descs = {
    'A': 'Gene: Gene symbol (HGNC)',
    'B': 'S1_Zscore: Mean Z-scored expression of Subtype 1 (S1) centroid',
    'C': 'S2_Zscore: Mean Z-scored expression of Subtype 2 (S2) centroid',
    'D': 'Training_Mean_log2TPM: Gene-wise mean log2(TPM+1) in TCGA-KIRC training set',
    'E': 'Training_SD_log2TPM: Gene-wise standard deviation of log2(TPM+1) in TCGA-KIRC training set',
}
for col, desc in descs.items():
    ws[f'{col}2'] = desc
    ws[f'{col}2'].font = Font(italic=True, size=9, color='555555')
    ws[f'{col}2'].alignment = Alignment(wrap_text=True)
ws.row_dimensions[2].height = 40

headers = ['Gene', 'S1_Zscore', 'S2_Zscore',
           'Training_Mean_log2TPM', 'Training_SD_log2TPM']
for i, h in enumerate(headers, 1):
    c = ws.cell(row=3, column=i, value=h)
    c.font = Font(bold=True, size=10)
    c.alignment = Alignment(horizontal='center')

for r_idx, (_, row) in enumerate(df.iterrows(), 4):
    ws.cell(row=r_idx, column=1, value=row['Gene'])
    for c_idx, col in enumerate(
            ['S1_Zscore', 'S2_Zscore',
             'Training_Mean_log2TPM', 'Training_SD_log2TPM'], 2):
        cell = ws.cell(row=r_idx, column=c_idx, value=row[col])
        cell.number_format = '0.0000'

ws.column_dimensions['A'].width = 14
for c in 'BCDE':
    ws.column_dimensions[c].width = 24

out = os.path.join(RESULTS, 'Table_S2_NCC_Classifier_68genes.xlsx')
wb.save(out)
print(f'Saved: {out} ({len(df)} genes)')